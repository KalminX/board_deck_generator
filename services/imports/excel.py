"""
Excel Ingestion Engine.
Reads and inspects multi-sheet .xlsx workbooks safely without macro execution.
Detects financial tables, column mappings, and extracts structured datasets.
"""
import openpyxl
from io import BytesIO
from typing import Dict, Any, List, Optional, Tuple
from decimal import Decimal
from services.finance.normalization import to_decimal


class ExcelImportService:
    """
    Parses canonical or semi-structured Excel financial workbooks.
    """

    MAX_FILE_SIZE_BYTES = 15 * 1024 * 1024  # 15 MB limit

    @staticmethod
    def inspect_workbook(file_or_bytes) -> Dict[str, Any]:
        """
        Quickly inspects sheet names and column headers without loading all data.
        """
        if isinstance(file_or_bytes, bytes):
            wb = openpyxl.load_workbook(filename=BytesIO(file_or_bytes), read_only=True, data_only=True)
        else:
            wb = openpyxl.load_workbook(filename=file_or_bytes, read_only=True, data_only=True)

        sheet_info = {}
        for sheet_name in wb.sheetnames:
            ws = wb[sheet_name]
            headers = []
            for row in ws.iter_rows(max_row=5, values_only=True):
                # find first non-empty row as header candidate
                non_empty = [str(c).strip() for c in row if c is not None and str(c).strip()]
                if len(non_empty) >= 2:
                    headers = [str(c).strip() if c is not None else "" for c in row]
                    break
            sheet_info[sheet_name] = {
                "detected_headers": headers,
                "max_rows": ws.max_row
            }
        wb.close()
        return sheet_info

    @classmethod
    def parse_workbook(cls, file_or_bytes) -> Dict[str, Any]:
        """
        Parses workbook sheets into normalized data dictionaries.
        Supports sheets:
        - Metadata / Summary
        - P&L / Income Statement
        - Balance Sheet
        - Cash Flow
        - KPIs
        - Actions & Risks
        """
        if isinstance(file_or_bytes, bytes):
            wb = openpyxl.load_workbook(filename=BytesIO(file_or_bytes), data_only=True)
        else:
            wb = openpyxl.load_workbook(filename=file_or_bytes, data_only=True)

        sheet_names_lower = {s.strip().lower(): s for s in wb.sheetnames}

        # 1. Parse Metadata / Overview if available
        meta = cls._parse_metadata_sheet(wb, sheet_names_lower)

        # 2. Parse Income Statement / P&L
        pnl_data = cls._parse_pnl_sheet(wb, sheet_names_lower)

        # 3. Parse Balance Sheet
        bs_data = cls._parse_balance_sheet(wb, sheet_names_lower)

        # 4. Parse Cash Flow
        cf_data = cls._parse_cash_flow_sheet(wb, sheet_names_lower)

        # 5. Parse KPIs
        kpis_data = cls._parse_kpis_sheet(wb, sheet_names_lower)

        # 6. Parse Risks & Actions
        risks_data, actions_data = cls._parse_risks_actions_sheet(wb, sheet_names_lower)

        wb.close()

        return {
            "organisation": meta.get("organisation", {}),
            "period": meta.get("period", {}),
            "income_statement": pnl_data,
            "balance_sheet": bs_data,
            "cash_flow": cf_data,
            "kpis": kpis_data,
            "risks": risks_data,
            "actions": actions_data,
            "metadata": meta.get("custom", {})
        }

    @classmethod
    def _parse_metadata_sheet(cls, wb, sheets: Dict[str, str]) -> Dict[str, Any]:
        meta_sheet_name = None
        for key in ["metadata", "summary", "overview", "cover", "info"]:
            if key in sheets:
                meta_sheet_name = sheets[key]
                break

        res = {
            "organisation": {"name": "Apex Technologies Holdings Ltd", "currency_code": "USD", "currency_symbol": "$"},
            "period": {"label": "Q3 2026", "period_type": "quarter", "prior_period_label": "Q2 2026", "budget_period_label": "Q3 2026 Budget"},
            "custom": {}
        }

        if not meta_sheet_name:
            return res

        ws = wb[meta_sheet_name]
        for row in ws.iter_rows(values_only=True):
            if not row or len(row) < 2 or row[0] is None:
                continue
            k = str(row[0]).strip().lower().replace(":", "").replace(" ", "_")
            v = str(row[1]).strip() if row[1] is not None else ""
            if not v:
                continue

            if k in ["company_name", "organisation", "company", "org"]:
                res["organisation"]["name"] = v
            elif k in ["currency", "currency_code"]:
                res["organisation"]["currency_code"] = v
                res["organisation"]["currency_symbol"] = "$" if v == "USD" else ("£" if v == "GBP" else ("€" if v == "EUR" else v))
            elif k in ["period", "reporting_period", "period_label"]:
                res["period"]["label"] = v
            elif k in ["prior_period", "prior_period_label"]:
                res["period"]["prior_period_label"] = v
            elif k in ["budget_period", "budget_label"]:
                res["period"]["budget_period_label"] = v
            else:
                res["custom"][k] = v

        return res

    @classmethod
    def _find_header_map(cls, ws) -> Tuple[int, Dict[str, int]]:
        """
        Finds header row index and column mapping for standard financial columns.
        """
        for r_idx, row in enumerate(ws.iter_rows(max_row=10, values_only=True), start=1):
            if not row:
                continue
            mapping = {}
            for c_idx, cell in enumerate(row):
                if cell is None:
                    continue
                txt = str(cell).strip().lower()
                if txt in ["line item", "account", "item", "description", "financial line", "metric", "account name", "property"]:
                    mapping["label"] = c_idx
                elif txt in ["actual", "actuals", "current", "current actual", "q3 actual", "month actual", "value"]:
                    mapping["actual"] = c_idx
                elif txt in ["budget", "target", "plan", "forecast"]:
                    mapping["budget"] = c_idx
                elif txt in ["prior", "prior period", "prior actual", "previous", "py"]:
                    mapping["prior"] = c_idx
                elif txt in ["ytd", "ytd actual", "year to date"]:
                    mapping["ytd_actual"] = c_idx
                elif txt in ["ytd budget", "ytd target"]:
                    mapping["ytd_budget"] = c_idx
                elif txt in ["section", "category", "type", "class", "bucket"]:
                    mapping["category"] = c_idx
                elif txt in ["notes", "commentary", "note"]:
                    mapping["notes"] = c_idx

            if "label" in mapping and ("actual" in mapping or "budget" in mapping or "prior" in mapping or "value" in mapping):
                return r_idx, mapping

        return 1, {"label": 0, "actual": 1, "budget": 2, "prior": 3, "ytd_actual": 4}

    @classmethod
    def _parse_pnl_sheet(cls, wb, sheets: Dict[str, str]) -> Dict[str, Any]:
        pnl_sheet_name = None
        for key in ["p&l", "pnl", "income statement", "profit and loss", "profit & loss"]:
            if key in sheets:
                pnl_sheet_name = sheets[key]
                break

        if not pnl_sheet_name:
            return {"revenue_items": [], "cogs_items": [], "opex_items": [], "other_items": []}

        ws = wb[pnl_sheet_name]
        header_row, col_map = cls._find_header_map(ws)

        revenue_items = []
        cogs_items = []
        opex_items = []
        other_items = []
        monthly_trend = []

        current_category = "revenue"

        for row in ws.iter_rows(min_row=header_row + 1, values_only=True):
            if not row or all(c is None for c in row):
                continue
            lbl_idx = col_map.get("label", 0)
            if lbl_idx >= len(row) or row[lbl_idx] is None:
                continue

            lbl = str(row[lbl_idx]).strip()
            if not lbl or lbl.startswith("#"):
                continue

            # Category detection from section headers
            lbl_lower = lbl.lower()
            if lbl_lower in ["revenue", "revenues", "income"]:
                current_category = "revenue"
                continue
            elif lbl_lower in ["cost of sales", "cogs", "cost of goods sold", "direct costs"]:
                current_category = "cogs"
                continue
            elif lbl_lower in ["operating expenses", "opex", "operating costs", "overheads"]:
                current_category = "opex"
                continue
            elif lbl_lower in ["other income / expense", "other expenses", "tax", "interest"]:
                current_category = "other"
                continue

            item_dict = {
                "key": lbl.lower().replace(" ", "_"),
                "label": lbl,
                "actual": row[col_map["actual"]] if "actual" in col_map and col_map["actual"] < len(row) else None,
                "budget": row[col_map["budget"]] if "budget" in col_map and col_map["budget"] < len(row) else None,
                "prior": row[col_map["prior"]] if "prior" in col_map and col_map["prior"] < len(row) else None,
                "ytd_actual": row[col_map["ytd_actual"]] if "ytd_actual" in col_map and col_map["ytd_actual"] < len(row) else None,
                "ytd_budget": row[col_map["ytd_budget"]] if "ytd_budget" in col_map and col_map["ytd_budget"] < len(row) else None,
                "category": current_category,
                "notes": str(row[col_map["notes"]]).strip() if "notes" in col_map and col_map["notes"] < len(row) and row[col_map["notes"]] is not None else None,
                "source_ref": f"{pnl_sheet_name}"
            }

            # Check if row is a total/subtotal
            if "total" in lbl_lower or "gross profit" in lbl_lower or "operating profit" in lbl_lower or "ebit" in lbl_lower:
                item_dict["is_subtotal"] = True

            # Group into bucket
            row_cat = str(row[col_map["category"]]).lower() if "category" in col_map and col_map["category"] < len(row) and row[col_map["category"]] else current_category
            if "rev" in row_cat:
                revenue_items.append(item_dict)
            elif "cog" in row_cat or "cost" in row_cat:
                cogs_items.append(item_dict)
            elif "opex" in row_cat or "operat" in row_cat or "admin" in row_cat or "sales" in row_cat or "marketing" in row_cat:
                opex_items.append(item_dict)
            else:
                other_items.append(item_dict)

        # Check for Trend Sheet if present
        trend_sheet_name = None
        for key in ["trend", "pnl_trend", "monthly_trend", "historical"]:
            if key in sheets:
                trend_sheet_name = sheets[key]
                break
        if trend_sheet_name:
            tws = wb[trend_sheet_name]
            for row in tws.iter_rows(min_row=2, values_only=True):
                if row and len(row) >= 3 and row[0] is not None:
                    rev_d = to_decimal(row[1])
                    opex_d = to_decimal(row[2]) if len(row) > 2 else None
                    profit_d = to_decimal(row[3]) if len(row) > 3 else None
                    monthly_trend.append({
                        "period": str(row[0]).strip(),
                        "revenue": str(rev_d) if rev_d is not None else None,
                        "opex": str(opex_d) if opex_d is not None else None,
                        "profit": str(profit_d) if profit_d is not None else None
                    })

        return {
            "revenue_items": revenue_items,
            "cogs_items": cogs_items,
            "opex_items": opex_items,
            "other_items": other_items,
            "monthly_trend": monthly_trend
        }

    @classmethod
    def _parse_balance_sheet(cls, wb, sheets: Dict[str, str]) -> Dict[str, Any]:
        bs_sheet_name = None
        for key in ["balance sheet", "bs", "statement of financial position"]:
            if key in sheets:
                bs_sheet_name = sheets[key]
                break

        if not bs_sheet_name:
            return {"current_assets": [], "non_current_assets": [], "current_liabilities": [], "non_current_liabilities": [], "equity_items": []}

        ws = wb[bs_sheet_name]
        header_row, col_map = cls._find_header_map(ws)

        current_assets = []
        non_current_assets = []
        current_liabilities = []
        non_current_liabilities = []
        equity_items = []

        current_cat = "current_assets"

        for row in ws.iter_rows(min_row=header_row + 1, values_only=True):
            if not row or all(c is None for c in row):
                continue
            lbl_idx = col_map.get("label", 0)
            if lbl_idx >= len(row) or row[lbl_idx] is None:
                continue

            lbl = str(row[lbl_idx]).strip()
            if not lbl or lbl.startswith("#"):
                continue

            lbl_lower = lbl.lower()
            if "current asset" in lbl_lower:
                current_cat = "current_assets"
                continue
            elif "non-current asset" in lbl_lower or "fixed asset" in lbl_lower or "non current asset" in lbl_lower:
                current_cat = "non_current_assets"
                continue
            elif "current liabilit" in lbl_lower:
                current_cat = "current_liabilities"
                continue
            elif "non-current liabilit" in lbl_lower or "long-term liabilit" in lbl_lower or "long term liabilit" in lbl_lower:
                current_cat = "non_current_liabilities"
                continue
            elif "equity" in lbl_lower or "shareholders" in lbl_lower:
                current_cat = "equity_items"
                continue

            item = {
                "key": lbl.lower().replace(" ", "_"),
                "label": lbl,
                "actual": row[col_map["actual"]] if "actual" in col_map and col_map["actual"] < len(row) else None,
                "prior": row[col_map["prior"]] if "prior" in col_map and col_map["prior"] < len(row) else None,
                "notes": str(row[col_map["notes"]]).strip() if "notes" in col_map and col_map["notes"] < len(row) and row[col_map["notes"]] is not None else None,
                "source_ref": bs_sheet_name
            }

            row_cat = str(row[col_map["category"]]).lower() if "category" in col_map and col_map["category"] < len(row) and row[col_map["category"]] else current_cat
            if "non-current asset" in row_cat or "fixed asset" in row_cat:
                non_current_assets.append(item)
            elif "current asset" in row_cat or "asset" in row_cat:
                current_assets.append(item)
            elif "non-current liab" in row_cat or "long-term liab" in row_cat or "long term liab" in row_cat:
                non_current_liabilities.append(item)
            elif "current liab" in row_cat or "liabilit" in row_cat:
                current_liabilities.append(item)
            elif "equity" in row_cat or "shareholder" in row_cat:
                equity_items.append(item)
            else:
                current_assets.append(item)

        return {
            "current_assets": current_assets,
            "non_current_assets": non_current_assets,
            "current_liabilities": current_liabilities,
            "non_current_liabilities": non_current_liabilities,
            "equity_items": equity_items
        }

    @classmethod
    def _parse_cash_flow_sheet(cls, wb, sheets: Dict[str, str]) -> Dict[str, Any]:
        cf_sheet_name = None
        for key in ["cash flow", "cashflow", "cf", "statement of cash flows"]:
            if key in sheets:
                cf_sheet_name = sheets[key]
                break

        if not cf_sheet_name:
            return {"operating_activities": [], "investing_activities": [], "financing_activities": []}

        ws = wb[cf_sheet_name]
        header_row, col_map = cls._find_header_map(ws)

        op = []
        inv = []
        fin = []
        current_cat = "operating"

        for row in ws.iter_rows(min_row=header_row + 1, values_only=True):
            if not row or all(c is None for c in row):
                continue
            lbl_idx = col_map.get("label", 0)
            if lbl_idx >= len(row) or row[lbl_idx] is None:
                continue

            lbl = str(row[lbl_idx]).strip()
            if not lbl or lbl.startswith("#"):
                continue

            lbl_lower = lbl.lower()
            if "operating" in lbl_lower:
                current_cat = "operating"
                continue
            elif "investing" in lbl_lower:
                current_cat = "investing"
                continue
            elif "financing" in lbl_lower:
                current_cat = "financing"
                continue

            item = {
                "key": lbl.lower().replace(" ", "_"),
                "label": lbl,
                "actual": row[col_map["actual"]] if "actual" in col_map and col_map["actual"] < len(row) else None,
                "budget": row[col_map["budget"]] if "budget" in col_map and col_map["budget"] < len(row) else None,
                "prior": row[col_map["prior"]] if "prior" in col_map and col_map["prior"] < len(row) else None,
                "source_ref": cf_sheet_name
            }

            if current_cat == "operating":
                op.append(item)
            elif current_cat == "investing":
                inv.append(item)
            elif current_cat == "financing":
                fin.append(item)

        return {
            "operating_activities": op,
            "investing_activities": inv,
            "financing_activities": fin
        }

    @classmethod
    def _parse_kpis_sheet(cls, wb, sheets: Dict[str, str]) -> List[Dict[str, Any]]:
        kpi_sheet_name = None
        for key in ["kpis", "kpi", "metrics", "key metrics"]:
            if key in sheets:
                kpi_sheet_name = sheets[key]
                break

        if not kpi_sheet_name:
            return []

        ws = wb[kpi_sheet_name]
        kpis = []
        for row in ws.iter_rows(min_row=2, values_only=True):
            if not row or row[0] is None or not str(row[0]).strip():
                continue
            kpis.append({
                "name": str(row[0]).strip(),
                "category": str(row[1]).strip() if len(row) > 1 and row[1] else "financial",
                "current_value": row[2] if len(row) > 2 else None,
                "target_value": row[3] if len(row) > 3 else None,
                "prior_value": row[4] if len(row) > 4 else None,
                "unit": str(row[5]).strip() if len(row) > 5 and row[5] else "$",
                "definition": str(row[6]).strip() if len(row) > 6 and row[6] else "",
                "higher_is_better": True if len(row) <= 7 or row[7] is None or str(row[7]).lower() in ("true", "1", "yes") else False
            })
        return kpis

    @classmethod
    def _parse_risks_actions_sheet(cls, wb, sheets: Dict[str, str]) -> Tuple[List[Dict[str, Any]], List[Dict[str, Any]]]:
        ra_sheet_name = None
        for key in ["risks & actions", "risks and actions", "risks_actions", "actions", "risks", "initiatives"]:
            if key in sheets:
                ra_sheet_name = sheets[key]
                break

        if not ra_sheet_name:
            return [], []

        ws = wb[ra_sheet_name]
        risks = []
        actions = []

        current_mode = "actions"

        for row in ws.iter_rows(min_row=1, values_only=True):
            if not row or all(c is None for c in row):
                continue
            first = str(row[0]).strip() if row[0] is not None else ""
            if not first:
                continue

            if first.lower() in ["risks", "risk items", "identified risks"]:
                current_mode = "risks"
                continue
            elif first.lower() in ["actions", "decisions", "priority actions", "action items"]:
                current_mode = "actions"
                continue
            elif first.lower() in ["title", "action", "risk", "item"]:
                continue

            if current_mode == "risks":
                risks.append({
                    "title": first,
                    "category": str(row[1]).strip() if len(row) > 1 and row[1] else "Operational",
                    "severity": str(row[2]).strip() if len(row) > 2 and row[2] else "Medium",
                    "likelihood": str(row[3]).strip() if len(row) > 3 and row[3] else "Medium",
                    "description": str(row[4]).strip() if len(row) > 4 and row[4] else "",
                    "mitigation_plan": str(row[5]).strip() if len(row) > 5 and row[5] else "",
                    "owner": str(row[6]).strip() if len(row) > 6 and row[6] else None
                })
            else:
                actions.append({
                    "title": first,
                    "reason": str(row[1]).strip() if len(row) > 1 and row[1] else "",
                    "priority": str(row[2]).strip() if len(row) > 2 and row[2] else "Medium",
                    "owner": str(row[3]).strip() if len(row) > 3 and row[3] else None,
                    "due_date": str(row[4]).strip() if len(row) > 4 and row[4] else None,
                    "status": str(row[5]).strip() if len(row) > 5 and row[5] else "Open"
                })

        return risks, actions
