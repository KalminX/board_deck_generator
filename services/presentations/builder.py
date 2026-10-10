"""
Presentation Builder Engine.
Orchestrates end-to-end PowerPoint generation from validated financial datasets.
Coordinates metric calculations, commentary generation, layout composition,
and native slide rendering.
"""
from io import BytesIO
from typing import List, Dict, Any, Optional, Tuple
from pptx import Presentation
from services.finance.canonical_schema import ReportContext
from services.finance.metrics import calculate_financial_metrics
from services.commentary.base import PresentationCommentary, CommentarySection
from services.commentary.factory import get_commentary_provider
from services.presentations.tokens import BrandConfig, DEFAULT_BRAND, MODERN_NAVY_BRAND, MINIMAL_EMERALD_BRAND
from services.presentations.validation import PresentationValidator, PresentationValidationResult

# Slide renderers
from services.presentations.slides.slide_01_cover import render_cover_slide
from services.presentations.slides.slide_02_executive_summary import render_executive_summary_slide
from services.presentations.slides.slide_03_financial_overview import render_financial_overview_slide
from services.presentations.slides.slide_04_pnl import render_pnl_slide
from services.presentations.slides.slide_05_balance_sheet import render_balance_sheet_slide
from services.presentations.slides.slide_06_cash_flow import render_cash_flow_slide
from services.presentations.slides.slide_07_kpis import render_kpi_slide
from services.presentations.slides.slide_08_risks import render_risks_slide
from services.presentations.slides.slide_09_actions import render_actions_slide
from services.presentations.slides.slide_10_appendix import render_appendix_slide


SLIDE_REGISTRY = {
    "cover": render_cover_slide,
    "executive_summary": render_executive_summary_slide,
    "financial_overview": render_financial_overview_slide,
    "pnl": render_pnl_slide,
    "balance_sheet": render_balance_sheet_slide,
    "cash_flow": render_cash_flow_slide,
    "kpis": render_kpi_slide,
    "risks": render_risks_slide,
    "actions": render_actions_slide,
    "appendix": render_appendix_slide,
}

SLIDE_TITLES = {
    "cover": "Cover Page",
    "executive_summary": "Executive Summary",
    "financial_overview": "Financial Performance Overview",
    "pnl": "Profit & Loss Statement",
    "balance_sheet": "Balance Sheet",
    "cash_flow": "Statement of Cash Flows",
    "kpis": "Key Performance Indicators",
    "risks": "Risks & Anomalies",
    "actions": "Actions & Decisions",
    "appendix": "Appendix & Governance",
}


class PresentationBuilder:
    """
    Main presentation generation orchestrator.
    """

    def __init__(self, brand: Optional[BrandConfig] = None, narrative_style: str = "standard", commentary_provider_name: Optional[str] = None):
        self.brand = brand or DEFAULT_BRAND
        self.narrative_style = narrative_style
        self.commentary_provider = get_commentary_provider(commentary_provider_name)

    def build_presentation(
        self,
        context: ReportContext,
        selected_slides: Optional[List[str]] = None,
        custom_commentary: Optional[Dict[str, Dict[str, Any]]] = None
    ) -> Tuple[Presentation, Dict[str, Any], PresentationValidationResult]:
        """
        Builds the PowerPoint presentation and returns (pptx_presentation_obj, manifest_dict, validation_res).
        """
        if selected_slides is None:
            selected_slides = list(SLIDE_REGISTRY.keys())

        # 1. Validate slide selection
        val_sel = PresentationValidator.validate_slide_selection(selected_slides)
        if not val_sel.is_valid:
            raise ValueError(f"Invalid slide selection: {', '.join(val_sel.errors)}")

        # 2. Compute Authoritative Metrics & Findings
        context = calculate_financial_metrics(context)

        # 3. Generate Commentary
        commentary = self.commentary_provider.generate_commentary(context, style=self.narrative_style)

        # 4. Apply custom commentary overrides (e.g. from user review edits)
        if custom_commentary:
            self._apply_commentary_overrides(commentary, custom_commentary)

        # 5. Initialize PowerPoint Presentation
        prs = Presentation()
        prs.slide_width = self.brand.slide_width
        prs.slide_height = self.brand.slide_height

        total_pages = len(selected_slides)
        rendered_slides = []

        # 6. Render Slides
        for page_idx, slide_key in enumerate(selected_slides, start=1):
            if slide_key not in SLIDE_REGISTRY:
                continue

            renderer = SLIDE_REGISTRY[slide_key]

            if slide_key == "cover":
                renderer(prs, context, self.brand)
            else:
                renderer(prs, context, commentary, current_page=page_idx, total_pages=total_pages, brand=self.brand)

            rendered_slides.append({
                "slide_number": page_idx,
                "slide_key": slide_key,
                "title": SLIDE_TITLES.get(slide_key, slide_key)
            })

        # 7. Validate Presentation Output
        val_out = PresentationValidator.validate_presentation_output(prs, expected_count=total_pages)

        # 8. Compile Generation Manifest
        manifest = {
            "organisation": context.organisation.name,
            "period": context.period.label,
            "brand_theme": self.brand.brand_name,
            "narrative_style": self.narrative_style,
            "slide_count": len(prs.slides),
            "slides": rendered_slides,
            "facts_count": len(context.all_facts),
            "findings_count": len(context.findings),
            "risks_count": len(context.risks),
            "actions_count": len(context.actions),
            "warnings": val_sel.warnings + val_out.warnings
        }

        return prs, manifest, val_out

    def _apply_commentary_overrides(self, commentary: PresentationCommentary, overrides: Dict[str, Dict[str, Any]]):
        for key, data in overrides.items():
            sec = commentary.get_section(key)
            if not sec:
                sec = CommentarySection(section_id=key, title=data.get("title", key.title()), lead_summary="", bullet_points=[])
                commentary.custom_sections[key] = sec

            if "lead_summary" in data:
                sec.lead_summary = str(data["lead_summary"])
            if "bullet_points" in data and isinstance(data["bullet_points"], list):
                sec.bullet_points = [str(b) for b in data["bullet_points"]]
            elif "bullets_text" in data:
                sec.bullet_points = [line.strip().lstrip("•-* ") for line in str(data["bullets_text"]).split("\n") if line.strip()]

    def save_to_file(self, prs: Presentation, filepath: str):
        prs.save(filepath)

    def save_to_bytes(self, prs: Presentation) -> BytesIO:
        bio = BytesIO()
        prs.save(bio)
        bio.seek(0)
        return bio


# Type alias for return signature clarity
Tuple_PRS_Manifest = tuple
