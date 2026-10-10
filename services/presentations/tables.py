"""
Native PowerPoint Table Builders.
Constructs formatted financial statements, risk matrices, and action item tables.
Enforces right-aligned decimal figures, header themes, and alternating row styling.
"""
from typing import List, Optional
from decimal import Decimal
from pptx.util import Inches, Pt
from pptx.enum.text import PP_ALIGN
from pptx.dml.color import RGBColor
from services.finance.canonical_schema import FinancialLineItem, RiskItem, ActionItem
from services.presentations.tokens import BrandConfig, DEFAULT_BRAND
from services.presentations.layouts import format_table_cell


def fmt_curr_table(val: Optional[Decimal], is_pct: bool = False) -> str:
    if val is None:
        return "—"
    if is_pct:
        return f"{val:+.1f}%" if val != 0 else "0.0%"
    
    # Financial formatting
    if val < 0:
        return f"({abs(val):,.0f})"
    elif val == 0:
        return "—"
    else:
        return f"{val:,.0f}"


def create_financial_statement_table(
    slide,
    left: Inches,
    top: Inches,
    width: Inches,
    height: Inches,
    items: List[FinancialLineItem],
    show_prior: bool = True,
    show_ytd: bool = True,
    currency_symbol: str = "$",
    brand: BrandConfig = DEFAULT_BRAND
):
    """
    Renders a native structured financial statement table.
    Columns: Line Item | Actual ($) | Budget ($) | Var ($) | Var (%) | [Prior] | [YTD]
    """
    headers = [f"Line Item ({currency_symbol})", "Actual", "Budget", "Var ($)", "Var (%)"]
    if show_prior:
        headers.append("Prior Period")
    if show_ytd:
        headers.append("YTD Actual")

    num_cols = len(headers)
    num_rows = len(items) + 1  # 1 header + rows

    table_shape = slide.shapes.add_table(num_rows, num_cols, left, top, width, height)
    table = table_shape.table

    # Set column widths
    table.columns[0].width = int(Inches(3.2))
    col_width = int((width - Inches(3.2)) / (num_cols - 1))
    for c in range(1, num_cols):
        table.columns[c].width = col_width

    # Render Header
    for c_idx, head in enumerate(headers):
        align = PP_ALIGN.LEFT if c_idx == 0 else PP_ALIGN.RIGHT
        format_table_cell(
            table.cell(0, c_idx),
            head,
            bold=True,
            font_size=brand.size_table_header,
            font_color=RGBColor(255, 255, 255),
            align=align,
            bg_color=brand.table_header_bg,
            font_name=brand.font_title
        )

    # Render Rows
    for r_idx, item in enumerate(items, start=1):
        is_bold = item.is_total or item.is_subtotal
        bg_col = brand.accent_light if item.is_total else (brand.table_alt_row if r_idx % 2 == 1 else RGBColor(255, 255, 255))
        txt_col = brand.primary_color if is_bold else brand.text_dark

        # Col 0: Label
        prefix = "    " if not is_bold and item.category != "general" else ""
        format_table_cell(table.cell(r_idx, 0), prefix + item.label, bold=is_bold, font_size=brand.size_table_cell, font_color=txt_col, align=PP_ALIGN.LEFT, bg_color=bg_col)

        # Col 1: Actual
        format_table_cell(table.cell(r_idx, 1), fmt_curr_table(item.actual), bold=is_bold, font_size=brand.size_table_cell, font_color=txt_col, align=PP_ALIGN.RIGHT, bg_color=bg_col)

        # Col 2: Budget
        format_table_cell(table.cell(r_idx, 2), fmt_curr_table(item.budget), bold=is_bold, font_size=brand.size_table_cell, font_color=txt_col, align=PP_ALIGN.RIGHT, bg_color=bg_col)

        # Col 3: Var ($)
        var_val = item.variance_vs_budget
        var_color = brand.success_color if (var_val or 0) > 0 else (brand.danger_color if (var_val or 0) < 0 else txt_col)
        format_table_cell(table.cell(r_idx, 3), fmt_curr_table(var_val), bold=is_bold, font_size=brand.size_table_cell, font_color=var_color, align=PP_ALIGN.RIGHT, bg_color=bg_col)

        # Col 4: Var (%)
        var_pct = item.variance_pct_vs_budget
        format_table_cell(table.cell(r_idx, 4), fmt_curr_table(var_pct, is_pct=True), bold=is_bold, font_size=brand.size_table_cell, font_color=var_color, align=PP_ALIGN.RIGHT, bg_color=bg_col)

        col_pos = 5
        if show_prior:
            format_table_cell(table.cell(r_idx, col_pos), fmt_curr_table(item.prior), bold=is_bold, font_size=brand.size_table_cell, font_color=txt_col, align=PP_ALIGN.RIGHT, bg_color=bg_col)
            col_pos += 1

        if show_ytd:
            format_table_cell(table.cell(r_idx, col_pos), fmt_curr_table(item.ytd_actual), bold=is_bold, font_size=brand.size_table_cell, font_color=txt_col, align=PP_ALIGN.RIGHT, bg_color=bg_col)

    return table_shape


def create_risks_table(
    slide,
    left: Inches,
    top: Inches,
    width: Inches,
    height: Inches,
    risks: List[RiskItem],
    brand: BrandConfig = DEFAULT_BRAND
):
    """
    Renders structured risk register table.
    """
    headers = ["Risk Item & Description", "Category", "Severity", "Mitigation Plan", "Owner"]
    num_cols = len(headers)
    num_rows = max(len(risks) + 1, 2)

    table_shape = slide.shapes.add_table(num_rows, num_cols, left, top, width, height)
    table = table_shape.table

    table.columns[0].width = int(Inches(3.5))
    table.columns[1].width = int(Inches(1.5))
    table.columns[2].width = int(Inches(1.2))
    table.columns[3].width = int(Inches(4.2))
    table.columns[4].width = int(Inches(1.7))

    for c_idx, head in enumerate(headers):
        format_table_cell(table.cell(0, c_idx), head, bold=True, font_size=brand.size_table_header, font_color=RGBColor(255, 255, 255), bg_color=brand.table_header_bg, font_name=brand.font_title)

    if not risks:
        format_table_cell(table.cell(1, 0), "No active high/medium risks recorded.", font_size=brand.size_table_cell)
        return table_shape

    for r_idx, risk in enumerate(risks, start=1):
        bg_col = brand.table_alt_row if r_idx % 2 == 1 else RGBColor(255, 255, 255)
        
        # Col 0: Title & Desc
        desc_text = f"{risk.title}\n{risk.description}" if risk.description else risk.title
        format_table_cell(table.cell(r_idx, 0), desc_text, font_size=Pt(8.5), bg_color=bg_col)
        format_table_cell(table.cell(r_idx, 1), risk.category, font_size=Pt(8.5), bg_color=bg_col)

        # Severity
        sev_color = brand.danger_color if risk.severity.lower() == "high" else (brand.warning_color if risk.severity.lower() == "medium" else brand.success_color)
        format_table_cell(table.cell(r_idx, 2), risk.severity.upper(), bold=True, font_size=Pt(8.5), font_color=sev_color, bg_color=bg_col)

        # Mitigation
        format_table_cell(table.cell(r_idx, 3), risk.mitigation_plan or "Active monitoring", font_size=Pt(8.5), bg_color=bg_col)

        # Owner
        format_table_cell(table.cell(r_idx, 4), risk.owner or "Executive Committee", font_size=Pt(8.5), bg_color=bg_col)

    return table_shape


def create_actions_table(
    slide,
    left: Inches,
    top: Inches,
    width: Inches,
    height: Inches,
    actions: List[ActionItem],
    brand: BrandConfig = DEFAULT_BRAND
):
    """
    Renders structured actions and decisions table.
    """
    headers = ["Action / Decision", "Rationale & Business Context", "Priority", "Owner", "Target Date", "Status"]
    num_cols = len(headers)
    num_rows = max(len(actions) + 1, 2)

    table_shape = slide.shapes.add_table(num_rows, num_cols, left, top, width, height)
    table = table_shape.table

    table.columns[0].width = int(Inches(3.2))
    table.columns[1].width = int(Inches(3.5))
    table.columns[2].width = int(Inches(1.1))
    table.columns[3].width = int(Inches(1.5))
    table.columns[4].width = int(Inches(1.3))
    table.columns[5].width = int(Inches(1.5))

    for c_idx, head in enumerate(headers):
        format_table_cell(table.cell(0, c_idx), head, bold=True, font_size=brand.size_table_header, font_color=RGBColor(255, 255, 255), bg_color=brand.table_header_bg, font_name=brand.font_title)

    if not actions:
        format_table_cell(table.cell(1, 0), "No active governance actions assigned.", font_size=brand.size_table_cell)
        return table_shape

    for r_idx, action in enumerate(actions, start=1):
        bg_col = brand.table_alt_row if r_idx % 2 == 1 else RGBColor(255, 255, 255)

        format_table_cell(table.cell(r_idx, 0), action.title, bold=True, font_size=Pt(8.5), bg_color=bg_col)
        format_table_cell(table.cell(r_idx, 1), action.reason or "Strategic roadmap milestone", font_size=Pt(8.5), bg_color=bg_col)

        p_color = brand.danger_color if action.priority.lower() in ("urgent", "high") else brand.text_dark
        format_table_cell(table.cell(r_idx, 2), action.priority.upper(), bold=True, font_size=Pt(8.5), font_color=p_color, bg_color=bg_col)

        format_table_cell(table.cell(r_idx, 3), action.owner or "Leadership Team", font_size=Pt(8.5), bg_color=bg_col)
        format_table_cell(table.cell(r_idx, 4), action.due_date or "Q4 2026", font_size=Pt(8.5), bg_color=bg_col)
        format_table_cell(table.cell(r_idx, 5), action.status.upper(), bold=True, font_size=Pt(8.5), bg_color=bg_col)

    return table_shape
