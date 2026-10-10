"""
Slide 7: Key Performance Indicators (KPIs) Analysis Slide.
"""
from pptx.util import Inches, Pt
from decimal import Decimal
from services.finance.canonical_schema import ReportContext
from services.commentary.base import PresentationCommentary
from services.presentations.tokens import BrandConfig, DEFAULT_BRAND
from services.presentations.layouts import apply_slide_header, apply_slide_footer, create_kpi_card, create_card_container
from services.presentations.charts import create_kpi_trend_chart
from services.commentary.templates import fmt_curr


def render_kpi_slide(prs, context: ReportContext, commentary: PresentationCommentary, current_page: int, total_pages: int, brand: BrandConfig = DEFAULT_BRAND):
    slide_layout = prs.slide_layouts[6]
    slide = prs.slides.add_slide(slide_layout)

    apply_slide_header(slide, "Key Performance Indicators & Scorecard", "Operational Performance", context.period.label, brand)
    apply_slide_footer(slide, current_page, total_pages, context.organisation.name, context.period.label, brand)

    kpis = context.kpis

    # Left Side: 2x3 Grid of KPI Metric Cards
    left_w = Inches(7.5)
    card_w = Inches(3.6)
    card_h = Inches(1.7)
    gap_x = Inches(0.25)
    gap_y = Inches(0.2)
    start_top = Inches(1.2)

    for idx, kpi in enumerate(kpis[:6]):
        row = idx // 2
        col = idx % 2
        x_pos = brand.margin_left + col * (card_w + gap_x)
        y_pos = start_top + row * (card_h + gap_y)

        # Format value
        val_str = f"{kpi.unit}{kpi.current_value:,.1f}" if kpi.unit == "$" and kpi.current_value is not None else (f"{kpi.current_value:,.1f}{kpi.unit}" if kpi.current_value is not None else "N/A")
        tgt_str = f"Target: {kpi.unit}{kpi.target_value:,.1f}" if kpi.target_value is not None else "Target: N/A"
        prior_str = f" | Prior: {kpi.prior_value:,.1f}" if kpi.prior_value is not None else ""
        sub_str = f"{tgt_str}{prior_str}\n{kpi.definition}" if kpi.definition else f"{tgt_str}{prior_str}"

        create_kpi_card(slide, x_pos, y_pos, card_w, card_h, kpi.name, val_str, sub_str, kpi.status, brand)

    # Right Side: KPI Trend Chart or Scorecard commentary
    right_left = brand.margin_left + left_w + Inches(0.25)
    right_w = Inches(4.38)

    # If first KPI has monthly history, chart it; otherwise show scorecard commentary
    has_chart = False
    if kpis and kpis[0].monthly_history:
        h = kpis[0].monthly_history
        cats = [d.get("period", "") for d in h]
        vals = [float(d.get("value", 0)) for d in h]
        tgts = [float(d.get("target", 0)) for d in h] if "target" in h[0] else None
        create_kpi_trend_chart(slide, right_left, Inches(1.2), right_w, Inches(2.7), kpis[0].name, cats, vals, tgts, kpis[0].unit, brand)
        has_chart = True

    c_top = Inches(4.1) if has_chart else Inches(1.2)
    c_h = Inches(2.7) if has_chart else Inches(5.6)

    # Scorecard commentary
    create_card_container(slide, right_left, c_top, right_w, c_h, "KPI PERFORMANCE SUMMARY", brand)
    sec = commentary.kpi_commentary
    if sec:
        tb = slide.shapes.add_textbox(right_left + Inches(0.18), c_top + Inches(0.42), right_w - Inches(0.36), c_h - Inches(0.5))
        tf = tb.text_frame
        tf.word_wrap = True
        tf.margin_left = tf.margin_top = tf.margin_right = tf.margin_bottom = 0

        p0 = tf.paragraphs[0]
        p0.text = sec.lead_summary
        p0.font.name = brand.font_body
        p0.font.size = Pt(9)
        p0.font.bold = True
        p0.font.color.rgb = brand.primary_color
        p0.space_after = Pt(4)

        for b in sec.bullet_points:
            pb = tf.add_paragraph()
            pb.text = f"•  {b}"
            pb.font.name = brand.font_body
            pb.font.size = Pt(8.5)
            pb.font.color.rgb = brand.secondary_color
            pb.space_after = Pt(3)

    return slide
