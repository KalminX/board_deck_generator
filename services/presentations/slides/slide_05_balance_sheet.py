"""
Slide 5: Balance Sheet Slide.
"""
from pptx.util import Inches, Pt
from decimal import Decimal
from services.finance.canonical_schema import ReportContext
from services.commentary.base import PresentationCommentary
from services.presentations.tokens import BrandConfig, DEFAULT_BRAND
from services.presentations.layouts import apply_slide_header, apply_slide_footer, create_commentary_box
from services.presentations.tables import create_financial_statement_table


def render_balance_sheet_slide(prs, context: ReportContext, commentary: PresentationCommentary, current_page: int, total_pages: int, brand: BrandConfig = DEFAULT_BRAND):
    slide_layout = prs.slide_layouts[6]
    slide = prs.slides.add_slide(slide_layout)

    apply_slide_header(slide, "Statement of Financial Position", "Balance Sheet", context.period.label, brand)
    apply_slide_footer(slide, current_page, total_pages, context.organisation.name, context.period.label, brand)

    bs = context.balance_sheet

    table_items = []
    # Current Assets
    for ca in bs.current_assets:
        table_items.append(ca)
    if bs.total_current_assets:
        table_items.append(bs.total_current_assets)

    # Non-Current Assets
    for nca in bs.non_current_assets:
        table_items.append(nca)
    if bs.total_assets:
        table_items.append(bs.total_assets)

    # Current Liabilities
    for cl in bs.current_liabilities:
        table_items.append(cl)
    if bs.total_current_liabilities:
        table_items.append(bs.total_current_liabilities)

    # Equity
    for eq in bs.equity_items:
        table_items.append(eq)
    if bs.total_equity:
        table_items.append(bs.total_equity)

    display_items = table_items[:12]

    table_w = Inches(8.3)
    table_h = Inches(5.6)

    create_financial_statement_table(
        slide,
        brand.margin_left, Inches(1.2), table_w, table_h,
        display_items,
        show_prior=True,
        show_ytd=False,
        currency_symbol=context.organisation.currency_symbol,
        brand=brand
    )

    # Right side commentary
    right_left = brand.margin_left + table_w + Inches(0.25)
    right_w = Inches(3.58)

    sec = commentary.balance_sheet_commentary
    create_commentary_box(slide, right_left, Inches(1.2), right_w, table_h, sec, "BALANCE SHEET ANALYSIS", brand)

    return slide
