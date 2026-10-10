"""
Slide 2: Executive Summary Slide.
"""
from pptx.util import Inches, Pt
from decimal import Decimal
from services.finance.canonical_schema import ReportContext
from services.commentary.base import PresentationCommentary
from services.presentations.tokens import BrandConfig, DEFAULT_BRAND
from services.presentations.layouts import apply_slide_header, apply_slide_footer, create_kpi_card, create_commentary_box, create_card_container
from services.commentary.templates import fmt_curr, fmt_pct


def render_executive_summary_slide(prs, context: ReportContext, commentary: PresentationCommentary, current_page: int, total_pages: int, brand: BrandConfig = DEFAULT_BRAND):
    slide_layout = prs.slide_layouts[6]
    slide = prs.slides.add_slide(slide_layout)

    apply_slide_header(slide, "Executive Summary & Key Findings", "Strategic Overview", context.period.label, brand)
    apply_slide_footer(slide, current_page, total_pages, context.organisation.name, context.period.label, brand)

    pnl = context.income_statement
    bs = context.balance_sheet
    sym = context.organisation.currency_symbol

    # Top Row: 4 Key Metric Cards
    card_w = Inches(2.9)
    card_h = Inches(1.15)
    top_pos = Inches(1.15)
    gap = Inches(0.18)

    # 1. Revenue
    rev_act = pnl.total_revenue.actual if pnl.total_revenue else None
    rev_var = pnl.total_revenue.variance_pct_vs_budget if pnl.total_revenue else None
    rev_sub = f"{fmt_pct(rev_var)} vs Budget" if rev_var is not None else "Actual Total"
    rev_status = "green" if (rev_var or 0) >= 0 else "red"
    create_kpi_card(slide, brand.margin_left, top_pos, card_w, card_h, "Total Revenue", fmt_curr(rev_act, sym), rev_sub, rev_status, brand)

    # 2. Gross Profit
    gp_act = pnl.gross_profit.actual if pnl.gross_profit else None
    gp_margin = (gp_act / rev_act * Decimal("100")) if gp_act and rev_act and rev_act > 0 else None
    gp_sub = f"{gp_margin:.1f}% Gross Margin" if gp_margin is not None else "Gross Margin"
    create_kpi_card(slide, brand.margin_left + (card_w + gap), top_pos, card_w, card_h, "Gross Profit", fmt_curr(gp_act, sym), gp_sub, "green" if (gp_margin or 0) >= 40 else "neutral", brand)

    # 3. Operating Profit (EBIT)
    op_act = pnl.operating_profit.actual if pnl.operating_profit else None
    op_var = pnl.operating_profit.variance_pct_vs_budget if pnl.operating_profit else None
    op_sub = f"{fmt_pct(op_var)} vs Budget" if op_var is not None else "Operating Profit"
    op_status = "green" if (op_var or 0) >= 0 else "red"
    create_kpi_card(slide, brand.margin_left + (card_w + gap) * 2, top_pos, card_w, card_h, "Operating Profit (EBIT)", fmt_curr(op_act, sym), op_sub, op_status, brand)

    # 4. Cash Balance
    cash_act = bs.cash_and_equivalents.actual if bs.cash_and_equivalents else None
    cash_sub = f"{context.cash_flow.cash_runway_months:.1f} mo runway" if context.cash_flow.cash_runway_months else "Ending Cash Reserves"
    create_kpi_card(slide, brand.margin_left + (card_w + gap) * 3, top_pos, card_w, card_h, "Cash Reserves", fmt_curr(cash_act, sym), cash_sub, "neutral", brand)

    # Bottom Area: 2-Column Split (Left: Executive Commentary, Right: Strategic Highlights & Priorities)
    bot_top = Inches(2.45)
    bot_h = Inches(4.45)
    col_w = Inches(5.95)

    # Left Column: Commentary
    sec = commentary.executive_summary
    create_commentary_box(slide, brand.margin_left, bot_top, col_w, bot_h, sec, "EXECUTIVE COMMENTARY & VARIANCE SUMMARY", brand)

    # Right Column: Strategic Highlights & Focus Areas Card
    create_card_container(slide, brand.margin_left + col_w + Inches(0.233), bot_top, col_w, bot_h, "STRATEGIC HIGHLIGHTS & OPERATIONAL FOCUS", brand)

    tb = slide.shapes.add_textbox(brand.margin_left + col_w + Inches(0.4), bot_top + Inches(0.45), col_w - Inches(0.35), bot_h - Inches(0.6))
    tf = tb.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_top = tf.margin_right = tf.margin_bottom = 0

    p1 = tf.paragraphs[0]
    p1.text = "Key Operational Takeaways:"
    p1.font.name = brand.font_title
    p1.font.size = Pt(10)
    p1.font.bold = True
    p1.font.color.rgb = brand.primary_color
    p1.space_after = Pt(4)

    # Highlight top variances and actions
    mat_findings = [f for f in context.findings if f.is_material]
    if mat_findings:
        for f in mat_findings[:2]:
            p = tf.add_paragraph()
            p.text = f"•  {f.headline} — {f.description}"
            p.font.name = brand.font_body
            p.font.size = Pt(8.5)
            p.font.color.rgb = brand.secondary_color
            p.space_after = Pt(3)

    p_act = tf.add_paragraph()
    p_act.text = "Board & Leadership Action Items:"
    p_act.font.name = brand.font_title
    p_act.font.size = Pt(10)
    p_act.font.bold = True
    p_act.font.color.rgb = brand.primary_color
    p_act.space_before = Pt(8)
    p_act.space_after = Pt(4)

    for a in context.actions[:3]:
        p = tf.add_paragraph()
        owner_str = f" ({a.owner})" if a.owner else ""
        p.text = f"•  [{a.priority.upper()}] {a.title}{owner_str}: {a.reason}"
        p.font.name = brand.font_body
        p.font.size = Pt(8.5)
        p.font.color.rgb = brand.secondary_color
        p.space_after = Pt(3)

    return slide
