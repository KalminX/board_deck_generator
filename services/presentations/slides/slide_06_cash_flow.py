"""
Slide 6: Statement of Cash Flows Slide.
"""
from pptx.util import Inches, Pt
from decimal import Decimal
from services.finance.canonical_schema import ReportContext
from services.commentary.base import PresentationCommentary
from services.presentations.tokens import BrandConfig, DEFAULT_BRAND
from services.presentations.layouts import apply_slide_header, apply_slide_footer, create_commentary_box
from services.presentations.tables import create_financial_statement_table


def render_cash_flow_slide(prs, context: ReportContext, commentary: PresentationCommentary, current_page: int, total_pages: int, brand: BrandConfig = DEFAULT_BRAND):
    slide_layout = prs.slide_layouts[6]
    slide = prs.slides.add_slide(slide_layout)

    apply_slide_header(slide, "Statement of Cash Flows", "Cash & Liquidity", context.period.label, brand)
    apply_slide_footer(slide, current_page, total_pages, context.organisation.name, context.period.label, brand)

    cf = context.cash_flow

    table_items = []
    # Operating
    for op in cf.operating_activities:
        table_items.append(op)
    if cf.net_cash_operating:
        table_items.append(cf.net_cash_operating)

    # Investing
    for inv in cf.investing_activities:
        table_items.append(inv)
    if cf.net_cash_investing:
        table_items.append(cf.net_cash_investing)

    # Financing
    for fin in cf.financing_activities:
        table_items.append(fin)
    if cf.net_cash_financing:
        table_items.append(cf.net_cash_financing)

    if cf.net_change_in_cash:
        table_items.append(cf.net_change_in_cash)

    display_items = table_items[:12]

    table_w = Inches(8.3)
    table_h = Inches(5.6)

    create_financial_statement_table(
        slide,
        brand.margin_left, Inches(1.2), table_w, table_h,
        display_items,
        show_prior=True,
        show_ytd=False,
        currency_symbol=context.organisation.currency_symbol,
        brand=brand
    )

    # Right side commentary
    right_left = brand.margin_left + table_w + Inches(0.25)
    right_w = Inches(3.58)

    sec = commentary.cash_flow_commentary
    create_commentary_box(slide, right_left, Inches(1.2), right_w, table_h, sec, "CASH & LIQUIDITY NOTES", brand)

    return slide
