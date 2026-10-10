"""
Slide 9: Actions and Decisions Slide.
"""
from pptx.util import Inches, Pt
from services.finance.canonical_schema import ReportContext
from services.commentary.base import PresentationCommentary
from services.presentations.tokens import BrandConfig, DEFAULT_BRAND
from services.presentations.layouts import apply_slide_header, apply_slide_footer, create_card_container
from services.presentations.tables import create_actions_table


def render_actions_slide(prs, context: ReportContext, commentary: PresentationCommentary, current_page: int, total_pages: int, brand: BrandConfig = DEFAULT_BRAND):
    slide_layout = prs.slide_layouts[6]
    slide = prs.slides.add_slide(slide_layout)

    apply_slide_header(slide, "Board Actions & Management Decisions", "Execution & Accountability", context.period.label, brand)
    apply_slide_footer(slide, current_page, total_pages, context.organisation.name, context.period.label, brand)

    # Top Context Summary
    top_w = Inches(12.133)
    top_h = Inches(1.1)

    create_card_container(slide, brand.margin_left, Inches(1.2), top_w, top_h, "EXECUTIVE MANDATES & GOVERNANCE COMMITMENTS", brand)

    sec = commentary.actions_commentary
    if sec:
        tb = slide.shapes.add_textbox(brand.margin_left + Inches(0.2), Inches(1.55), top_w - Inches(0.4), top_h - Inches(0.4))
        tf = tb.text_frame
        tf.word_wrap = True
        tf.margin_left = tf.margin_top = tf.margin_right = tf.margin_bottom = 0

        p = tf.paragraphs[0]
        p.text = f"{sec.lead_summary} " + " ".join(sec.bullet_points)
        p.font.name = brand.font_body
        p.font.size = Pt(9)
        p.font.color.rgb = brand.primary_color

    # Table Area
    bot_top = Inches(2.45)
    bot_h = Inches(4.35)

    create_actions_table(
        slide,
        brand.margin_left, bot_top, top_w, bot_h,
        context.actions[:8],
        brand=brand
    )

    return slide
