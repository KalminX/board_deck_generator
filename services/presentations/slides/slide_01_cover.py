"""
Slide 1: Cover Presentation Slide.
"""
from pptx.util import Inches, Pt
from pptx.enum.text import PP_ALIGN
from pptx.enum.shapes import MSO_SHAPE
from pptx.dml.color import RGBColor
from datetime import datetime
from services.finance.canonical_schema import ReportContext
from services.presentations.tokens import BrandConfig, DEFAULT_BRAND


def render_cover_slide(prs, context: ReportContext, brand: BrandConfig = DEFAULT_BRAND):
    slide_layout = prs.slide_layouts[6]  # Blank layout
    slide = prs.slides.add_slide(slide_layout)

    # Primary Dark Background for Cover
    bg = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, prs.slide_width, prs.slide_height)
    bg.fill.solid()
    bg.fill.fore_color.rgb = brand.primary_color
    bg.line.fill.background()

    # Left Decorative Accent Bar
    accent_bar = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.8), Inches(1.8), Inches(0.12), Inches(3.8))
    accent_bar.fill.solid()
    accent_bar.fill.fore_color.rgb = brand.accent_color
    accent_bar.line.fill.background()

    # Company Logo / Badge Placeholder
    badge = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(1.2), Inches(1.8), Inches(3.2), Inches(0.45))
    badge.fill.solid()
    badge.fill.fore_color.rgb = brand.secondary_color
    badge.line.fill.background()
    p_badge = badge.text_frame.paragraphs[0]
    p_badge.text = context.organisation.name.upper()
    p_badge.font.name = brand.font_title
    p_badge.font.size = Pt(9.5)
    p_badge.font.bold = True
    p_badge.font.color.rgb = RGBColor(255, 255, 255)
    p_badge.alignment = PP_ALIGN.CENTER

    # Main Title Box
    title_box = slide.shapes.add_textbox(Inches(1.2), Inches(2.4), Inches(10.5), Inches(1.8))
    tf_t = title_box.text_frame
    tf_t.word_wrap = True
    tf_t.margin_left = tf_t.margin_top = tf_t.margin_right = tf_t.margin_bottom = 0

    p_main = tf_t.paragraphs[0]
    p_main.text = "BOARD OF DIRECTORS REPORT"
    p_main.font.name = brand.font_title
    p_main.font.size = Pt(32)
    p_main.font.bold = True
    p_main.font.color.rgb = RGBColor(255, 255, 255)

    p_sub = tf_t.add_paragraph()
    p_sub.text = f"Financial Review & Management Accounts  |  {context.period.label}"
    p_sub.font.name = brand.font_title
    p_sub.font.size = Pt(18)
    p_sub.font.color.rgb = brand.card_border
    p_sub.space_before = Pt(8)

    # Metadata Card (Bottom Left)
    meta_card = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(1.2), Inches(4.5), Inches(7.5), Inches(1.5))
    meta_card.fill.solid()
    meta_card.fill.fore_color.rgb = brand.secondary_color
    meta_card.line.color.rgb = brand.accent_color
    meta_card.line.width = Pt(1)

    tf_m = meta_card.text_frame
    tf_m.word_wrap = True
    tf_m.margin_left = Inches(0.25)
    tf_m.margin_top = Inches(0.2)

    p_m1 = tf_m.paragraphs[0]
    p_m1.text = f"REPORTING PERIOD: {context.period.label}"
    p_m1.font.name = brand.font_title
    p_m1.font.size = Pt(10)
    p_m1.font.bold = True
    p_m1.font.color.rgb = RGBColor(255, 255, 255)

    p_m2 = tf_m.add_paragraph()
    p_m2.text = f"CURRENCY: {context.organisation.currency_code} ({context.organisation.currency_symbol})"
    p_m2.font.name = brand.font_body
    p_m2.font.size = Pt(9.5)
    p_m2.font.color.rgb = brand.card_border

    p_m3 = tf_m.add_paragraph()
    pub_date = datetime.now().strftime("%B %d, %Y")
    p_m3.text = f"PUBLISHED: {pub_date}   |   STATUS: Confidential Board Material"
    p_m3.font.name = brand.font_body
    p_m3.font.size = Pt(9.5)
    p_m3.font.color.rgb = brand.card_border

    return slide
