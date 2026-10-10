"""
Slide 8: Risks and Anomalies Slide.
"""
from pptx.util import Inches, Pt
from services.finance.canonical_schema import ReportContext
from services.commentary.base import PresentationCommentary
from services.presentations.tokens import BrandConfig, DEFAULT_BRAND
from services.presentations.layouts import apply_slide_header, apply_slide_footer, create_card_container
from services.presentations.tables import create_risks_table


def render_risks_slide(prs, context: ReportContext, commentary: PresentationCommentary, current_page: int, total_pages: int, brand: BrandConfig = DEFAULT_BRAND):
    slide_layout = prs.slide_layouts[6]
    slide = prs.slides.add_slide(slide_layout)

    apply_slide_header(slide, "Risks, Anomalies & Material Variances", "Risk Governance", context.period.label, brand)
    apply_slide_footer(slide, current_page, total_pages, context.organisation.name, context.period.label, brand)

    # Top Section: Material Variances & Anomalies Callout
    top_w = Inches(12.133)
    top_h = Inches(1.8)

    create_card_container(slide, brand.margin_left, Inches(1.2), top_w, top_h, "MATERIAL VARIANCE FINDINGS & AUDIT FLAGS", brand)

    tb = slide.shapes.add_textbox(brand.margin_left + Inches(0.2), Inches(1.6), top_w - Inches(0.4), top_h - Inches(0.5))
    tf = tb.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_top = tf.margin_right = tf.margin_bottom = 0

    mat_findings = [f for f in context.findings if f.is_material]
    if mat_findings:
        for idx, f in enumerate(mat_findings[:3]):
            p = tf.paragraphs[0] if idx == 0 else tf.add_paragraph()
            status_flag = "[INVESTIGATION REQUIRED]" if f.requires_investigation else "[FLAGGED]"
            p.text = f"•  {status_flag} {f.headline}: {f.description}"
            p.font.name = brand.font_body
            p.font.size = Pt(8.5)
            p.font.color.rgb = brand.danger_color if f.requires_investigation else brand.secondary_color
            p.space_after = Pt(2)
    else:
        p = tf.paragraphs[0]
        p.text = "•  All income statement and balance sheet lines remained within the standard 5.0% variance tolerance thresholds."
        p.font.name = brand.font_body
        p.font.size = Pt(8.5)
        p.font.color.rgb = brand.text_muted

    # Bottom Section: Risk Register Table
    bot_top = Inches(3.15)
    bot_h = Inches(3.65)

    create_risks_table(
        slide,
        brand.margin_left, bot_top, top_w, bot_h,
        context.risks[:6],
        brand=brand
    )

    return slide
