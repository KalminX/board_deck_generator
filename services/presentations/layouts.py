"""
Slide Layout Helpers and UI Components.
Constructs consistent corporate slide headers, footers, cards, commentary blocks,
and badges using python-pptx native shapes and text frames.
"""
from typing import Optional, List
from pptx.util import Inches, Pt
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE
from pptx.dml.color import RGBColor
from services.presentations.tokens import BrandConfig, DEFAULT_BRAND
from services.commentary.base import CommentarySection


def apply_slide_header(slide, title: str, category_badge: str, period_label: str, brand: BrandConfig = DEFAULT_BRAND):
    """
    Renders top banner, category badge, slide title, and period tag.
    """
    # Top accent line
    top_line = slide.shapes.add_shape(
        MSO_SHAPE.RECTANGLE,
        brand.margin_left, Inches(0.4), Inches(12.133), Inches(0.04)
    )
    top_line.fill.solid()
    top_line.fill.fore_color.rgb = brand.accent_color
    top_line.line.fill.background()

    # Category Badge (e.g., 'EXECUTIVE OVERVIEW' or 'FINANCIAL STATEMENTS')
    badge_box = slide.shapes.add_textbox(brand.margin_left, Inches(0.48), Inches(4.0), Inches(0.25))
    tf_b = badge_box.text_frame
    tf_b.word_wrap = True
    tf_b.margin_left = tf_b.margin_top = tf_b.margin_right = tf_b.margin_bottom = 0
    p_b = tf_b.paragraphs[0]
    p_b.text = category_badge.upper()
    p_b.font.name = brand.font_title
    p_b.font.size = Pt(8.5)
    p_b.font.bold = True
    p_b.font.color.rgb = brand.accent_color

    # Slide Title
    title_box = slide.shapes.add_textbox(brand.margin_left, Inches(0.72), Inches(8.5), Inches(0.45))
    tf_t = title_box.text_frame
    tf_t.word_wrap = True
    tf_t.margin_left = tf_t.margin_top = tf_t.margin_right = tf_t.margin_bottom = 0
    p_t = tf_t.paragraphs[0]
    p_t.text = title
    p_t.font.name = brand.font_title
    p_t.font.size = brand.size_slide_title
    p_t.font.bold = True
    p_t.font.color.rgb = brand.primary_color

    # Reporting Period Tag (Right-aligned)
    period_box = slide.shapes.add_textbox(Inches(9.5), Inches(0.65), Inches(3.2), Inches(0.35))
    tf_p = period_box.text_frame
    tf_p.word_wrap = True
    tf_p.margin_left = tf_p.margin_top = tf_p.margin_right = tf_p.margin_bottom = 0
    p_p = tf_p.paragraphs[0]
    p_p.alignment = PP_ALIGN.RIGHT
    p_p.text = f"PERIOD: {period_label.upper()}"
    p_p.font.name = brand.font_title
    p_p.font.size = Pt(9)
    p_p.font.bold = True
    p_p.font.color.rgb = brand.text_muted


def apply_slide_footer(slide, current_page: int, total_pages: int, company_name: str, period_label: str, brand: BrandConfig = DEFAULT_BRAND):
    """
    Renders standard bottom footer with confidentiality notice and page number.
    """
    footer_top = Inches(7.05)
    
    # Subtle divider
    div = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, brand.margin_left, footer_top, Inches(12.133), Inches(0.01))
    div.fill.solid()
    div.fill.fore_color.rgb = brand.card_border
    div.line.fill.background()

    # Left note: Confidential + Company
    left_box = slide.shapes.add_textbox(brand.margin_left, footer_top + Inches(0.05), Inches(6.0), Inches(0.3))
    tf_l = left_box.text_frame
    tf_l.word_wrap = True
    tf_l.margin_left = tf_l.margin_top = tf_l.margin_right = tf_l.margin_bottom = 0
    p_l = tf_l.paragraphs[0]
    p_l.text = f"CONFIDENTIAL  |  {company_name}  —  Management & Board Report"
    p_l.font.name = brand.font_body
    p_l.font.size = Pt(8)
    p_l.font.color.rgb = brand.text_muted

    # Right note: Page number
    right_box = slide.shapes.add_textbox(Inches(10.5), footer_top + Inches(0.05), Inches(2.233), Inches(0.3))
    tf_r = right_box.text_frame
    tf_r.word_wrap = True
    tf_r.margin_left = tf_r.margin_top = tf_r.margin_right = tf_r.margin_bottom = 0
    p_r = tf_r.paragraphs[0]
    p_r.alignment = PP_ALIGN.RIGHT
    p_r.text = f"Slide {current_page} of {total_pages}"
    p_r.font.name = brand.font_body
    p_r.font.size = Pt(8)
    p_r.font.color.rgb = brand.text_muted


def create_card_container(slide, left: Inches, top: Inches, width: Inches, height: Inches, title: Optional[str] = None, brand: BrandConfig = DEFAULT_BRAND):
    """
    Creates a clean background card container with border and optional card title.
    """
    card = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top, width, height)
    card.fill.solid()
    card.fill.fore_color.rgb = brand.card_bg
    card.line.color.rgb = brand.card_border
    card.line.width = Pt(1)

    if title:
        tbox = slide.shapes.add_textbox(left + Inches(0.18), top + Inches(0.12), width - Inches(0.36), Inches(0.3))
        tf = tbox.text_frame
        tf.word_wrap = True
        tf.margin_left = tf.margin_top = tf.margin_right = tf.margin_bottom = 0
        p = tf.paragraphs[0]
        p.text = title.upper()
        p.font.name = brand.font_title
        p.font.size = brand.size_card_title
        p.font.bold = True
        p.font.color.rgb = brand.secondary_color

    return card


def create_kpi_card(
    slide,
    left: Inches,
    top: Inches,
    width: Inches,
    height: Inches,
    label: str,
    value_str: str,
    subtext: str,
    status: str = "neutral",
    brand: BrandConfig = DEFAULT_BRAND
):
    """
    Renders an executive KPI summary card.
    """
    # Background card
    card = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top, width, height)
    card.fill.solid()
    card.fill.fore_color.rgb = brand.card_bg
    card.line.color.rgb = brand.card_border
    card.line.width = Pt(1)

    # Status indicator left bar
    status_bar = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top, Inches(0.08), height)
    status_bar.fill.solid()
    if status == "green":
        status_bar.fill.fore_color.rgb = brand.success_color
    elif status == "red":
        status_bar.fill.fore_color.rgb = brand.danger_color
    elif status == "amber":
        status_bar.fill.fore_color.rgb = brand.warning_color
    else:
        status_bar.fill.fore_color.rgb = brand.accent_color
    status_bar.line.fill.background()

    # Text frame
    tb = slide.shapes.add_textbox(left + Inches(0.18), top + Inches(0.12), width - Inches(0.25), height - Inches(0.24))
    tf = tb.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_top = tf.margin_right = tf.margin_bottom = 0

    # Metric Label
    p_lbl = tf.paragraphs[0]
    p_lbl.text = label.upper()
    p_lbl.font.name = brand.font_title
    p_lbl.font.size = brand.size_metric_label
    p_lbl.font.bold = True
    p_lbl.font.color.rgb = brand.text_muted

    # Metric Value
    p_val = tf.add_paragraph()
    p_val.text = value_str
    p_val.font.name = brand.font_title
    p_val.font.size = brand.size_metric_value
    p_val.font.bold = True
    p_val.font.color.rgb = brand.primary_color

    # Subtext (vs Budget / Prior)
    p_sub = tf.add_paragraph()
    p_sub.text = subtext
    p_sub.font.name = brand.font_body
    p_sub.font.size = Pt(8.5)
    if status == "green":
        p_sub.font.color.rgb = brand.success_color
    elif status == "red":
        p_sub.font.color.rgb = brand.danger_color
    else:
        p_sub.font.color.rgb = brand.text_muted


def create_commentary_box(
    slide,
    left: Inches,
    top: Inches,
    width: Inches,
    height: Inches,
    section: Optional[CommentarySection],
    header_title: str = "EXECUTIVE COMMENTARY & ANALYSIS",
    brand: BrandConfig = DEFAULT_BRAND
):
    """
    Renders structured commentary card with lead paragraph and bullet points.
    """
    # Outer card
    create_card_container(slide, left, top, width, height, header_title, brand)

    if not section:
        return

    # Text content
    tb = slide.shapes.add_textbox(left + Inches(0.2), top + Inches(0.42), width - Inches(0.4), height - Inches(0.5))
    tf = tb.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_top = tf.margin_right = tf.margin_bottom = 0

    # Lead summary
    p_lead = tf.paragraphs[0]
    p_lead.text = section.lead_summary
    p_lead.font.name = brand.font_body
    p_lead.font.size = Pt(9.5)
    p_lead.font.bold = True
    p_lead.font.color.rgb = brand.primary_color
    p_lead.space_after = Pt(6)

    # Bullet points
    for bullet in section.bullet_points:
        p_b = tf.add_paragraph()
        p_b.text = f"•  {bullet}"
        p_b.font.name = brand.font_body
        p_b.font.size = Pt(9)
        p_b.font.color.rgb = brand.secondary_color
        p_b.space_after = Pt(4)


def format_table_cell(
    cell,
    text: str,
    bold: bool = False,
    font_size: Pt = Pt(9),
    font_color: RGBColor = RGBColor(15, 23, 42),
    align: PP_ALIGN = PP_ALIGN.LEFT,
    bg_color: Optional[RGBColor] = None,
    font_name: str = "Calibri"
):
    """
    Applies standard formatting to a PowerPoint table cell.
    """
    if bg_color:
        cell.fill.solid()
        cell.fill.fore_color.rgb = bg_color
    else:
        cell.fill.solid()
        cell.fill.fore_color.rgb = RGBColor(255, 255, 255)

    cell.vertical_anchor = MSO_ANCHOR.MIDDLE
    cell.margin_left = Inches(0.08)
    cell.margin_right = Inches(0.08)
    cell.margin_top = Inches(0.04)
    cell.margin_bottom = Inches(0.04)

    tf = cell.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0

    p = tf.paragraphs[0]
    p.text = str(text)
    p.alignment = align
    p.font.name = font_name
    p.font.size = font_size
    p.font.bold = bold
    p.font.color.rgb = font_color
